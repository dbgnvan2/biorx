"""
BioRxiv Research Tool - PyQt6 GUI
Main application with Search & Browse and Filters tabs.
"""

import sys
import json
import logging
import re
import threading
from pathlib import Path
from typing import Optional, Dict, Any, List
from datetime import datetime, timedelta

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QTabWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QListWidget, QListWidgetItem, QCheckBox,
    QSpinBox, QDateEdit, QComboBox, QTextEdit, QSplitter, QMessageBox,
    QProgressBar, QTableWidget, QTableWidgetItem, QHeaderView, QFrame,
    QFormLayout, QGroupBox, QDialog, QInputDialog, QScrollArea, QPlainTextEdit,
)
from PyQt6.QtCore import Qt, QDate, pyqtSignal, QThread, QObject
from PyQt6.QtGui import QFont

from agents.search_agent import SearchAgent
from agents.summarization_agent import SummarizationAgent
from src.db import Database
from src.biorxiv_api import BioRxivAPI
from src.pdf_handler import PDFHandler
from src.sources.config import (
    load_sources_config, get_enabled_search_sources,
    get_default_selected_sources, SOURCE_LABELS,
)
from src.sources.orchestrator import SourceOrchestrator
from src.selection import ResultsSelection, ResultsAccumulator, paper_key

_LOG_FILE = Path("biorx.log")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(_LOG_FILE, encoding="utf-8"),
    ],
)
logger = logging.getLogger(__name__)

FILTERS_PATH = Path("filters.json")

BIORXIV_CATEGORIES = [
    "(any)",
    "animal behavior and cognition",
    "biochemistry",
    "bioengineering",
    "bioinformatics",
    "biophysics",
    "cancer biology",
    "cell biology",
    "clinical trials",
    "developmental biology",
    "ecology",
    "epidemiology",
    "evolutionary biology",
    "genetics",
    "genomics",
    "immunology",
    "microbiology",
    "molecular biology",
    "neuroscience",
    "paleontology",
    "pathology",
    "pharmacology and toxicology",
    "physiology",
    "plant biology",
    "scientific communication and education",
    "synthetic biology",
    "systems biology",
    "zoology",
]

PAPER_TYPES = ["(any)", "new results", "confirmatory results", "contradictory results", "review article"]
VERSIONS    = ["(any)", "1 (first submission only)", "2+ (revised only)"]
PUBLISHED   = ["(any)", "preprints only (not in journal)", "published in journal only"]
LICENSES    = ["(any)", "cc_by", "cc_by_nc", "cc_by_nd", "cc_no", "pd"]
SPECIES     = [
    "(any)",
    "Human studies only",
    "Exclude animal studies",
    "Animal studies only",
]


from src.filters_store import (     # noqa: E402  (kept at their former home)
    load_filters_file as load_filters,
    save_filters_file as save_filters,
    filter_is_enabled,
    filter_has_text as _filter_has_text,
)


def filter_initial_check_state() -> "Qt.CheckState":
    """Purpose: Decide the startup check state of a saved-filter row in the Search panel.
    Spec:    docs/implementation_plan_2026-06-07.md#E1.1
    Tests:   tests/test_gui_filters.py::test_e1_1_filters_unchecked_on_startup

    Saved-filter rows always start unchecked so a run only happens on an
    explicit selection, regardless of the filter's persisted ``enabled`` flag.
    """
    return Qt.CheckState.Unchecked


# ---------------------------------------------------------------------------
# Background workers
# ---------------------------------------------------------------------------

class SearchWorker(QObject):
    batch_ready = pyqtSignal(list)      # matched papers from one API page
    progress    = pyqtSignal(int, int)  # (fetched so far, total)
    status      = pyqtSignal(str)       # terminal status (Done / Stopped)
    phase       = pyqtSignal(str)       # live phase text (per-source / enriching)
    # [(paper_key of the streamed row, enriched dict)] once enrichment has run.
    # Keyed, not by identity: a list signal hands the GUI copies of the dicts.
    refreshed   = pyqtSignal(object)
    finished    = pyqtSignal(list)      # all matched papers
    error       = pyqtSignal(str)

    MAX_PAPERS = 2000

    def __init__(self, orchestrator: "SourceOrchestrator", filter_dict: Dict[str, Any],
                 save_to_db: bool = False, db: Optional[Database] = None):
        super().__init__()
        self.orchestrator = orchestrator
        self.filter_dict  = filter_dict
        self.save_to_db   = save_to_db
        self.db           = db
        self._stop_event  = threading.Event()
        self._all_matched: List[Dict] = []

    def stop(self):
        self._stop_event.set()

    def run(self):
        try:
            f = self.filter_dict
            self._all_matched = []
            source_selection  = f.get("source_selection", {"all": True, "selected": []})
            # Records that passed the filter; only these are enriched
            # (plan 2026-09-18 C2). Rows are streamed as pages arrive, before
            # enrichment; each streamed dict is paired with its record so the
            # rows can be refreshed once enrichment has run (issue 4).
            matched_ids: set = set()
            pairs: List[Any] = []     # (record, streamed dict)

            def on_batch(records):
                if self._stop_event.is_set():
                    return
                papers  = [r.to_dict() for r in records]
                matched = _filter_papers(papers, f)
                kept = {id(p) for p in matched}
                for r, p in zip(records, papers):
                    if id(p) in kept:
                        matched_ids.add(id(r))
                        pairs.append((r, p))
                if matched:
                    self._all_matched.extend(matched)
                    self.batch_ready.emit(matched)
                    # Live status is composed UI-side in _append_batch so it can
                    # show both this filter's count and the running total.

            def on_progress(fetched: int, total: int):
                self.progress.emit(fetched, max(fetched, total))

            def on_status(message: str):
                if not self._stop_event.is_set():
                    self.phase.emit(message)

            self.orchestrator.search(
                filter_dict=f,
                source_selection=source_selection,
                on_batch=on_batch,
                on_progress=on_progress,
                on_status=on_status,
                should_stop=self._stop_event.is_set,
                max_results=self.MAX_PAPERS,
                enrich_only=lambda r: id(r) in matched_ids,
            )

            # What enrichment added (PDF links, licence, filled abstracts)
            # reaches the table and the saved rows (issue 4; learnings P36).
            updates = [(paper_key(streamed), record.to_dict()) for record, streamed in pairs]
            self._all_matched = [fresh for _, fresh in updates]
            if updates:
                self.refreshed.emit(updates)

            if self._stop_event.is_set():
                self.status.emit(f"Stopped — {len(self._all_matched):,} matched")
            elif self.save_to_db and self.db:
                saved = sum(1 for p in self._all_matched if self.db.insert_paper(p))
                self.status.emit(
                    f"Done — {len(self._all_matched):,} matched, {saved:,} saved"
                )

            self.finished.emit(self._all_matched)
        except Exception as e:
            self.error.emit(str(e))
        finally:
            # This worker runs on its own QThread and took a database
            # connection on it. Give it back now rather than waiting for the
            # thread to be torn down (src/db.py — Database.release).
            if self.db:
                self.db.release()


class SummarizationWorker(QObject):
    finished = pyqtSignal()
    error    = pyqtSignal(str)
    progress = pyqtSignal(str)

    def __init__(self, agent: SummarizationAgent, paper_id: Optional[int] = None, max_count: int = 10):
        super().__init__()
        self.agent = agent
        self.paper_id = paper_id
        self.max_count = max_count

    def run(self):
        try:
            self.progress.emit("Starting summarization…")
            if self.paper_id:
                ok = self.agent.summarize_paper_by_id(self.paper_id)
                msg = f"Summary generated for paper {self.paper_id}" if ok else f"Failed to summarize paper {self.paper_id}"
                self.progress.emit(msg)
            else:
                result = self.agent.summarize_all_unsummarized(max_count=self.max_count)
                self.progress.emit(f"Complete: {result.get('summarized_count', 0)} papers summarized")
            self.finished.emit()
        except Exception as e:
            self.error.emit(str(e))


# ---------------------------------------------------------------------------
# Paper detail dialog + download worker
# ---------------------------------------------------------------------------

from src.paper_meta import (        # noqa: E402  (kept at their former home)
    recover_abstract,
    pdf_url as _pdf_url,
    paper_link as _paper_link,
    scrape_abstract_from_url as _scrape_abstract_from_url,
    fetch_openalex_abstract as _fetch_openalex_abstract,
)


def _open_in_browser(url: str):
    import webbrowser
    webbrowser.open(url)


class AbstractFetchWorker(QObject):
    """
    Fetch a missing abstract in a background thread.

    The strategy chain lives in src/paper_meta.recover_abstract so the web app
    uses the same one (docs/implementation_plan_2026-09-16_backlog.md#N2).
    This worker only moves it off the GUI thread and turns a miss into a
    display message — the message is shown, never passed on as an abstract.
    """
    finished = pyqtSignal(str)
    error    = pyqtSignal(str)

    NOT_FOUND_MESSAGE = "(Abstract not available from any source)"

    def __init__(self, paper: Dict[str, Any]):
        super().__init__()
        self.paper = dict(paper)

    def run(self):
        try:
            result = recover_abstract(self.paper)
            if result.found:
                logger.debug("AbstractFetchWorker: recovered abstract via %s for doi=%s",
                             result.source, self.paper.get("doi", "(no doi)"))
                self.finished.emit(result.text)
            else:
                logger.warning("AbstractFetchWorker: no abstract found for doi=%s title=%r",
                               self.paper.get("doi"), self.paper.get("title", "")[:80])
                self.finished.emit(self.NOT_FOUND_MESSAGE)
        except Exception as e:
            logger.error("AbstractFetchWorker: exception for doi=%s: %s", self.paper.get("doi"), e)
            self.error.emit(str(e))


class BatchPdfDownloadWorker(QObject):
    """Download a list of papers' PDFs sequentially in a background thread."""
    progress = pyqtSignal(int, int, str)   # (done, total, current_title)
    finished = pyqtSignal(int, int)        # (succeeded, failed)
    stopped  = pyqtSignal()

    def __init__(self, papers: list):
        super().__init__()
        self.papers   = papers
        self._stop    = False

    def stop(self):
        self._stop = True

    def run(self):
        from src.pdf_handler import PDFHandler
        handler   = PDFHandler()
        total     = len(self.papers)
        succeeded = 0
        failed    = 0
        for i, paper in enumerate(self.papers):
            if self._stop:
                self.stopped.emit()
                return
            title = paper.get("title", "Untitled")[:60]
            self.progress.emit(i, total, title)
            url = _pdf_url(paper)
            doi = paper.get("doi", "")
            if not url:
                logger.warning("Batch PDF download: no URL for %r (doi=%s)", title, doi)
                failed += 1
                continue
            try:
                path = handler.download_pdf(url, paper.get("title", ""), doi)
                if path:
                    succeeded += 1
                else:
                    failed += 1
            except Exception as e:
                logger.warning("Batch PDF download failed for %s: %s", doi, e)
                failed += 1
        self.finished.emit(succeeded, failed)


class PdfSectionWorker(QObject):
    """Extract sections from a local PDF in a background thread."""
    finished = pyqtSignal(str)   # discussion text
    error    = pyqtSignal(str)

    def __init__(self, pdf_path: str):
        super().__init__()
        self.pdf_path = pdf_path

    def run(self):
        from src.pdf_handler import PDFHandler
        try:
            handler    = PDFHandler()
            sections   = handler.extract_sections(self.pdf_path)
            discussion = sections.get("discussion", "").strip()
            if not discussion:
                discussion = "(Discussion section not found — the PDF layout may use non-standard headings)"
            self.finished.emit(discussion)
        except Exception as e:
            self.error.emit(str(e))


class PaperDetailDialog(QDialog):
    """Show abstract and discussion for a paper."""

    def __init__(self, paper: Dict[str, Any], parent=None):
        super().__init__(parent)
        self.paper   = paper
        self._thread = None
        self._worker = None
        self.setWindowTitle("Paper Detail")
        self.setMinimumSize(860, 640)
        self.init_ui()

    def init_ui(self):
        from PyQt6.QtWidgets import QTabWidget, QFileDialog
        layout = QVBoxLayout()

        # ── Header ──
        title_label = QLabel(self.paper.get("title", ""))
        title_label.setWordWrap(True)
        title_label.setStyleSheet("font-size: 14pt; font-weight: bold;")
        layout.addWidget(title_label)

        authors = self.paper.get("authors", "")
        if isinstance(authors, list):
            authors = "; ".join(authors)
        meta = (
            f"<b>Authors:</b> {authors}<br>"
            f"<b>Date:</b> {self.paper.get('date') or self.paper.get('pub_date', '')}  "
            f"&nbsp;|&nbsp; <b>Category:</b> {self.paper.get('category', '')}  "
            f"&nbsp;|&nbsp; <b>Type:</b> {self.paper.get('type', '')}  "
            f"&nbsp;|&nbsp; <b>DOI:</b> {self.paper.get('doi', '')}"
        )
        meta_label = QLabel(meta)
        meta_label.setWordWrap(True)
        layout.addWidget(meta_label)

        line = QFrame(); line.setFrameShape(QFrame.Shape.HLine)
        layout.addWidget(line)

        # ── Tabs: Abstract | Discussion ──
        tabs = QTabWidget()

        # Abstract tab
        abstract_widget = QWidget()
        al = QVBoxLayout()
        self.abstract_text = QTextEdit()
        self.abstract_text.setReadOnly(True)
        abstract = self.paper.get("abstract") or ""
        lookup_keys = ("doi", "pmcid", "best_oa_url", "source_url", "url")
        if abstract:
            self.abstract_text.setPlainText(abstract)
        elif any(self.paper.get(k) for k in lookup_keys):
            self.abstract_text.setPlainText("Retrieving abstract…")
            self._fetch_abstract()
        else:
            self.abstract_text.setPlainText("(No abstract available)")
        al.addWidget(self.abstract_text)
        abstract_widget.setLayout(al)
        tabs.addTab(abstract_widget, "Abstract")

        # Discussion tab
        discussion_widget = QWidget()
        dl = QVBoxLayout()
        self.discussion_text = QTextEdit()
        self.discussion_text.setReadOnly(True)
        self.discussion_text.setPlainText(
            "To view the discussion section:\n\n"
            "1. Click 'Open PDF in Browser' — your browser will download the PDF.\n"
            "2. Once downloaded, click 'Load from local PDF…' and select the file.\n\n"
            "The discussion section will be extracted and shown here."
        )
        dl.addWidget(self.discussion_text)

        disc_btn_row = QHBoxLayout()
        open_browser_btn = QPushButton("Open PDF in Browser")
        open_browser_btn.clicked.connect(lambda: _open_in_browser(_pdf_url(self.paper)))
        self.load_local_btn = QPushButton("Load from local PDF…")
        self.load_local_btn.clicked.connect(self._pick_and_extract)
        self.disc_status = QLabel("")
        disc_btn_row.addWidget(open_browser_btn)
        disc_btn_row.addWidget(self.load_local_btn)
        disc_btn_row.addStretch()
        dl.addLayout(disc_btn_row)
        dl.addWidget(self.disc_status)
        discussion_widget.setLayout(dl)
        tabs.addTab(discussion_widget, "Discussion")

        layout.addWidget(tabs, 1)

        # ── Footer ──
        btn_row = QHBoxLayout()
        open_page_btn = QPushButton("Open Paper Page in Browser")
        page_url = (
            self.paper.get("source_url")
            or self.paper.get("url")
            or (f"https://doi.org/{self.paper['doi']}" if self.paper.get("doi") else "")
        )
        open_page_btn.clicked.connect(lambda u=page_url: _open_in_browser(u) if u else None)
        close_btn = QPushButton("Close")
        close_btn.clicked.connect(self.accept)
        btn_row.addWidget(open_page_btn)
        btn_row.addStretch()
        btn_row.addWidget(close_btn)
        layout.addLayout(btn_row)

        self.setLayout(layout)

    def _fetch_abstract(self):
        """Fetch the abstract in the background; update the text widget when done."""
        self._abstract_thread = QThread()
        self._abstract_worker = AbstractFetchWorker(self.paper)
        self._abstract_worker.moveToThread(self._abstract_thread)
        self._abstract_thread.started.connect(self._abstract_worker.run)
        self._abstract_worker.finished.connect(self.abstract_text.setPlainText)
        self._abstract_worker.finished.connect(self._abstract_thread.quit)
        self._abstract_worker.error.connect(
            lambda e: self.abstract_text.setPlainText(f"(Could not fetch abstract: {e})")
        )
        self._abstract_worker.error.connect(self._abstract_thread.quit)
        self._abstract_thread.start()

    def _pick_and_extract(self):
        from PyQt6.QtWidgets import QFileDialog
        path, _ = QFileDialog.getOpenFileName(self, "Select PDF", str(Path.home()), "PDF files (*.pdf)")
        if not path:
            return
        self.load_local_btn.setEnabled(False)
        self.disc_status.setText("Extracting…")
        self.discussion_text.setPlainText("Extracting discussion section…")

        self._worker = PdfSectionWorker(path)
        self._thread = QThread()
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.finished.connect(self._on_extracted)
        self._worker.error.connect(self._on_extract_error)
        self._thread.start()

    def _on_extracted(self, discussion: str):
        self.discussion_text.setPlainText(discussion)
        self.disc_status.setText("Done.")
        self.load_local_btn.setEnabled(True)

    def _on_extract_error(self, err: str):
        self.disc_status.setText(f"Error: {err}")
        self.load_local_btn.setEnabled(True)


def _attach_context_menu(table: "QTableWidget"):
    """Add right-click context menu to a results table. Paper dict must be stored in col-0 UserRole."""
    table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)

    def on_context_menu(pos):
        from PyQt6.QtWidgets import QMenu
        row = table.rowAt(pos.y())
        if row < 0:
            return
        item = table.item(row, 0)
        if not item:
            return
        paper = item.data(Qt.ItemDataRole.UserRole)
        if not paper:
            return
        # SavedReferencesTab stores {"paper": {...}, "item_id": N}; unwrap.
        if isinstance(paper, dict) and "paper" in paper and "item_id" in paper:
            paper = paper["paper"]

        menu = QMenu(table)
        view_action     = menu.addAction("View Abstract & Discussion")
        browser_action  = menu.addAction("Open PDF in Browser")
        page_action     = menu.addAction("Open Paper Page in Browser")

        action = menu.exec(table.viewport().mapToGlobal(pos))
        if action == view_action:
            dlg = PaperDetailDialog(paper, table)
            dlg.exec()
        elif action == browser_action:
            _open_in_browser(_pdf_url(paper))
        elif action == page_action:
            url = (
                paper.get("source_url")
                or paper.get("url")
                or (f"https://doi.org/{paper['doi']}" if paper.get("doi") else "")
            )
            if url:
                _open_in_browser(url)

    table.customContextMenuRequested.connect(on_context_menu)


# ---------------------------------------------------------------------------
# Shared filtering logic (used by both tabs)
#
# The implementation lives in src/filtering.py so the headless CLI
# (agents/monitor.py) applies exactly the same filter semantics as the GUI.
# These module-level aliases keep the existing call sites unchanged.
# ---------------------------------------------------------------------------

from src.filtering import (          # noqa: E402  (kept next to its former home)
    split_terms        as _terms,
    match_term         as _match,
    text_group_matches as _text_group_matches,
    filter_papers      as _filter_papers,
)


# ---------------------------------------------------------------------------
# Source picker widget
# ---------------------------------------------------------------------------

class SourcePickerWidget(QWidget):
    """
    Compact source selection control.
    Shows 'All Sources' + one checkbox per enabled search source.
    Crossref / Unpaywall are enrichment-only and never appear here.
    """
    selection_changed = pyqtSignal(dict)  # {"all": bool, "selected": list[str]}

    def __init__(self, available_sources: List[str], default_selected: List[str],
                 parent=None):
        super().__init__(parent)
        self._available  = available_sources
        self._updating   = False

        layout = QHBoxLayout()
        layout.setContentsMargins(0, 0, 0, 0)

        self._all_cb = QCheckBox("All Sources")
        self._all_cb.setChecked(True)
        self._all_cb.stateChanged.connect(self._on_all_changed)
        layout.addWidget(self._all_cb)

        self._source_cbs: Dict[str, QCheckBox] = {}
        for source in available_sources:
            label = SOURCE_LABELS.get(source, source)
            cb = QCheckBox(label)
            cb.setChecked(source in default_selected)
            cb.setEnabled(not self._all_cb.isChecked())
            cb.stateChanged.connect(lambda state, s=source: self._on_source_changed(s, state))
            self._source_cbs[source] = cb
            layout.addWidget(cb)

        layout.addStretch()
        self.setLayout(layout)

    def _on_all_changed(self, state: int):
        if self._updating:
            return
        is_checked = state == Qt.CheckState.Checked.value
        self._updating = True
        for cb in self._source_cbs.values():
            cb.setEnabled(not is_checked)
            if is_checked:
                cb.setChecked(False)
        self._updating = False
        self.selection_changed.emit(self.get_selection())

    def _on_source_changed(self, source: str, state: int):
        if self._updating:
            return
        self._updating = True
        if state == Qt.CheckState.Checked.value:
            # At least one individual source selected → uncheck All
            self._all_cb.setChecked(False)
        else:
            # If no individual sources remain checked, restore All
            if not any(cb.isChecked() for cb in self._source_cbs.values()):
                self._all_cb.setChecked(True)
        self._updating = False
        self.selection_changed.emit(self.get_selection())

    def get_selection(self) -> Dict[str, Any]:
        if self._all_cb.isChecked():
            return {"all": True, "selected": []}
        selected = [s for s, cb in self._source_cbs.items() if cb.isChecked()]
        return {"all": False, "selected": selected}

    def load_selection(self, sel: Dict[str, Any]):
        self._updating = True
        use_all = sel.get("all", True)
        selected = sel.get("selected", [])
        self._all_cb.setChecked(use_all)
        for source, cb in self._source_cbs.items():
            cb.setChecked(not use_all and source in selected)
            cb.setEnabled(not use_all)
        self._updating = False


# ---------------------------------------------------------------------------
# Search & Browse tab
# ---------------------------------------------------------------------------

class SearchBrowseTab(QWidget):
    reference_list_saved = pyqtSignal()   # emitted after a new reference list is saved

    def __init__(self, db: Database, orchestrator: "SourceOrchestrator"):
        super().__init__()
        self.db           = db
        self.orchestrator = orchestrator
        self.summ_agent   = SummarizationAgent()
        self._results = ResultsAccumulator()  # dedups results across filters
        self.current_results: List[Dict[str, Any]] = self._results.papers
        self.current_page   = 0
        self.results_per_page = 20
        self._selection = ResultsSelection()  # page-independent paper selection (F1)
        self._threads: list = []
        # Source picker (reuses orchestrator's enabled sources)
        enabled  = orchestrator.get_enabled_sources()
        defaults = [s for s in enabled]  # all enabled = default for quick search
        self._source_picker = SourcePickerWidget(enabled, defaults)
        self.init_ui()

    def init_ui(self):
        layout = QHBoxLayout()

        # ── Left panel ──────────────────────────────────────────────────────
        left = QWidget()
        ll   = QVBoxLayout()

        ll.addWidget(QLabel("Saved Filters:"))
        self.filters_list = QListWidget()
        self.filters_list.setMinimumWidth(200)
        ll.addWidget(self.filters_list, 3)
        self.load_filters_list()

        run_layout = QHBoxLayout()
        self.run_selected_btn = QPushButton("Run Selected")
        self.run_selected_btn.clicked.connect(self.run_selected)
        self.run_all_btn = QPushButton("Run All Enabled")
        self.run_all_btn.clicked.connect(self.run_all_enabled)
        run_layout.addWidget(self.run_selected_btn)
        run_layout.addWidget(self.run_all_btn)
        ll.addLayout(run_layout)

        sep = QFrame(); sep.setFrameShape(QFrame.Shape.HLine); ll.addWidget(sep)

        ll.addWidget(QLabel("Quick Search:"))
        ll.addWidget(QLabel("From:"))
        self.from_date = QDateEdit(); self.from_date.setDate(QDate.currentDate().addDays(-7))
        ll.addWidget(self.from_date)
        ll.addWidget(QLabel("To:"))
        self.to_date = QDateEdit(); self.to_date.setDate(QDate.currentDate())
        ll.addWidget(self.to_date)
        ll.addWidget(QLabel("Category:"))
        self.cat_combo = QComboBox(); self.cat_combo.addItems(BIORXIV_CATEGORIES)
        ll.addWidget(self.cat_combo)
        ll.addWidget(QLabel("Keywords (comma-sep):"))
        self.quick_keywords = QLineEdit()
        ll.addWidget(self.quick_keywords)

        ll.addWidget(QLabel("Sources:"))
        ll.addWidget(self._source_picker)

        quick_btn = QPushButton("Search")
        quick_btn.clicked.connect(self.quick_search)
        ll.addWidget(quick_btn)

        ll.addStretch()
        left.setLayout(ll)

        # ── Right panel ─────────────────────────────────────────────────────
        right  = QWidget()
        rl     = QVBoxLayout()

        self.status_label = QLabel("Ready")
        rl.addWidget(self.status_label)

        prog_row = QHBoxLayout()
        self.progress_bar = QProgressBar()
        self.progress_bar.setTextVisible(True)
        self.progress_bar.setFormat("Fetched %v / %m papers")
        self.progress_bar.setVisible(False)
        self.stop_btn = QPushButton("⏹  Stop")
        self.stop_btn.setVisible(False)
        self.stop_btn.clicked.connect(self._stop_search)
        prog_row.addWidget(self.progress_bar)
        prog_row.addWidget(self.stop_btn)
        rl.addLayout(prog_row)

        self.results_table = QTableWidget()
        self.results_table.setColumnCount(6)
        self.results_table.setHorizontalHeaderLabels(["Title", "Authors", "Date", "Category", "Type", "Source"])
        self.results_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.results_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.results_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.results_table.setWordWrap(True)
        self.results_table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.results_table.setToolTip("Right-click a row to view abstract/discussion or download PDF")
        self.results_table.itemChanged.connect(self._on_item_changed)
        _attach_context_menu(self.results_table)
        rl.addWidget(self.results_table, 5)

        # ── Save-selected bar ────────────────────────────────────────────────
        save_bar = QHBoxLayout()
        sel_all_btn = QPushButton("☑ Select All"); sel_all_btn.clicked.connect(self._select_all)
        sel_none_btn = QPushButton("☐ Clear");     sel_none_btn.clicked.connect(self._select_none)
        self._checked_label = QLabel("0 selected")
        self._save_refs_btn = QPushButton("📁  Save as Reference List…")
        self._save_refs_btn.clicked.connect(self._save_selected_as_reference)
        save_bar.addWidget(sel_all_btn)
        save_bar.addWidget(sel_none_btn)
        save_bar.addWidget(self._checked_label)
        save_bar.addStretch()
        save_bar.addWidget(self._save_refs_btn)
        rl.addLayout(save_bar)

        page_layout = QHBoxLayout()
        self.prev_btn = QPushButton("← Prev"); self.prev_btn.clicked.connect(self.prev_page)
        self.next_btn = QPushButton("Next →"); self.next_btn.clicked.connect(self.next_page)
        self.page_label = QLabel("Page 1")
        page_layout.addWidget(self.prev_btn)
        page_layout.addWidget(self.page_label)
        page_layout.addWidget(self.next_btn)
        page_layout.addStretch()
        self.matches_label = QLabel("")
        self.matches_label.setToolTip(
            "Total Matches counts every hit across all run filters; "
            "Unduplicated is the unique set (what gets saved)."
        )
        page_layout.addWidget(self.matches_label)
        rl.addLayout(page_layout)

        right.setLayout(rl)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(left)
        splitter.addWidget(right)
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 3)
        layout.addWidget(splitter)
        self.setLayout(layout)

    def load_filters_list(self):
        self.filters_list.clear()
        for f in load_filters(FILTERS_PATH):
            item = QListWidgetItem(f["name"])
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(filter_initial_check_state())
            item.setData(Qt.ItemDataRole.UserRole, f)
            self.filters_list.addItem(item)

    def run_selected(self):
        filters = [
            self.filters_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.filters_list.count())
            if self.filters_list.item(i).checkState() == Qt.CheckState.Checked
        ]
        if not filters:
            QMessageBox.warning(self, "Nothing selected", "Check at least one filter")
            return
        self._run_filters(filters, save_to_db=True)

    def run_all_enabled(self):
        filters = [
            self.filters_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.filters_list.count())
            if filter_is_enabled(self.filters_list.item(i).data(Qt.ItemDataRole.UserRole))
        ]
        if not filters:
            QMessageBox.information(self, "No enabled filters", "Enable at least one filter in the Filters tab")
            return
        self._run_filters(filters, save_to_db=True)

    def _run_filters(self, filters: List[Dict], save_to_db: bool = False):
        self.run_selected_btn.setEnabled(False)
        self.run_all_btn.setEnabled(False)
        self._results.reset()    # clears current_results in place + match counts
        self._selection.clear()  # a new search starts with nothing selected (F1)
        self.matches_label.setText("")
        self.results_table.setRowCount(0)
        self.current_page = 0
        self._filter_queue  = list(filters)
        self._save_to_db    = save_to_db
        self._total_filters = len(filters)
        self._filter_idx    = 0
        self._current_filter_name    = ""
        self._current_filter_matched = 0
        self._current_phase          = ""
        self._progress_format        = "Fetched %v / %m papers"
        self._current_worker: Optional[SearchWorker] = None
        # Indeterminate bar while we don't know total yet
        self.progress_bar.setMaximum(0)
        self.progress_bar.setFormat("Searching…")
        self.progress_bar.setVisible(True)
        self.stop_btn.setVisible(True)
        self.status_label.setStyleSheet("font-weight: bold; color: #0055aa;")
        self.status_label.setText("⏳  Search started…")
        self._run_next_filter()

    def _stop_search(self):
        self._filter_queue.clear()
        if self._current_worker:
            self._current_worker.stop()
        self.stop_btn.setEnabled(False)
        self.status_label.setText("Stopping…")

    def _run_next_filter(self):
        if not self._filter_queue:
            self._on_all_filters_done()
            return
        f = self._filter_queue.pop(0)
        self._filter_idx += 1
        name = f.get("name", "")
        self._current_filter_name    = name
        self._current_filter_matched = 0
        self._current_phase          = ""
        # Skip filters with no search terms — they'd return an unfiltered date scan
        if not _filter_has_text(f):
            self.status_label.setText(
                f"⚠  Skipped '{name}': no search terms configured. "
                "Add keywords in the Text Search Groups."
            )
            self._run_next_filter()
            return
        self.status_label.setText(self._search_status_text(phase="searching…"))
        # Stay indeterminate until first progress signal
        self.progress_bar.setMaximum(0)
        self.progress_bar.setFormat("Connecting…")

        worker = SearchWorker(self.orchestrator, f, save_to_db=self._save_to_db, db=self.db)
        self._current_worker = worker
        thread = QThread()
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.batch_ready.connect(self._append_batch)
        worker.refreshed.connect(self._apply_enrichment)
        worker.progress.connect(self._update_progress)
        worker.phase.connect(self._on_phase)
        worker.status.connect(self.status_label.setText)
        worker.error.connect(lambda e: self.status_label.setText(f"⚠  {e}"))
        worker.finished.connect(lambda _: self._run_next_filter())
        thread.start()
        self._threads.append((thread, worker))

    def _search_status_text(self, phase: str = "") -> str:
        """Compose the top status line: the live phase (current source being
        queried, or the enrichment step) plus the running total across all
        filters in the run. Falls back to the per-filter matched count when no
        live phase is active."""
        total = len(self.current_results)
        matched = self._current_filter_matched
        live = self._current_phase or phase
        body = live if live else f"{matched:,} matched"
        if self._total_filters > 1:
            prefix = f"⏳  Filter {self._filter_idx}/{self._total_filters} '{self._current_filter_name}'"
            return f"{prefix}: {body} · {total:,} total so far"
        if live:
            # Quick / single search: lead with the live phase rather than the
            # (often unhelpful, e.g. "quick") filter name.
            return f"⏳  {body} · {total:,} so far"
        label = f"'{self._current_filter_name}'" if self._current_filter_name else "Search"
        return f"⏳  {label}: {body}"

    def _on_phase(self, message: str):
        """Slot for SearchWorker.phase — the current source/enrichment phase."""
        self._current_phase = message
        if message.startswith("Enriching"):
            self._progress_format = "Enriching %v / %m papers"
        else:
            self._progress_format = "Fetched %v / %m papers"
        self.status_label.setText(self._search_status_text())

    def _update_progress(self, fetched: int, total: int):
        if total > 0:
            self.progress_bar.setMaximum(total)
            self.progress_bar.setFormat(self._progress_format)
            self.progress_bar.setValue(fetched)

    def _append_batch(self, papers: list):
        # Dedup across filters: only genuinely-new papers join current_results,
        # while total_matches tracks every hit (shown as Total vs Unduplicated).
        self._results.add_batch(papers)
        self._current_filter_matched += len(papers)
        # Render through the single paginated path so live results and page
        # navigation behave identically (B1).
        self.display_page()
        self.status_label.setText(self._search_status_text())
        self._update_matches_label()
        self._update_checked_count()

    def _apply_enrichment(self, updates: list):
        """Purpose: Refresh streamed rows with what enrichment added.
        Spec:    docs/implementation_plan_2026-09-18_filter_run.md#I4
        Tests:   tests/test_gui_filters.py::test_i4_gui_rows_get_enriched_fields

        Runs on the GUI thread (a queued signal). Rows are matched by the
        paper_key they were streamed with and updated in place, so
        current_results and the selection model keep their references.
        """
        rows = {paper_key(p): p for p in self.current_results}
        for key, fresh in updates:
            row = rows.get(key)
            if row is not None:
                row.update(fresh)
        self.display_page()

    def _update_matches_label(self):
        total  = self._results.total_matches
        unique = self._results.unique_count
        if total == unique:
            self.matches_label.setText(f"Matches: {total:,}")
        else:
            self.matches_label.setText(
                f"Total Matches: {total:,}    Unduplicated Matches: {unique:,}"
            )

    def _on_all_filters_done(self):
        self._current_phase = ""
        # Top the bar out so the completed run reads as 100% before it hides.
        if self.progress_bar.maximum() > 0:
            self.progress_bar.setValue(self.progress_bar.maximum())
        self.progress_bar.setVisible(False)
        self.stop_btn.setVisible(False)
        self.stop_btn.setEnabled(True)
        self.status_label.setStyleSheet("")
        self.status_label.setText(
            f"✅  Done — {self._results.unique_count:,} unique papers found"
        )
        self._update_matches_label()
        self.run_selected_btn.setEnabled(True)
        self.run_all_btn.setEnabled(True)

    def quick_search(self):
        category = self.cat_combo.currentText()
        keywords = [k.strip() for k in self.quick_keywords.text().split(",") if k.strip()]
        days_back = self.from_date.date().daysTo(self.to_date.date()) or 7
        f = {
            "name": "quick",
            "category": category,
            "days_back": days_back,
            "keywords": keywords,
            "text_groups": [{"title": "", "abstract": "", "both": ", ".join(keywords)}] if keywords else [],
            "authors": [],
            "institution": "",
            "paper_type": "(any)",
            "version": "(any)",
            "published": "(any)",
            "license": "(any)",
            "source_selection": self._source_picker.get_selection(),
        }
        self.status_label.setText("Searching…")
        self._run_filters([f], save_to_db=False)

    def display_page(self):
        """Render the current page of current_results. Single rendering path for
        both live search and Prev/Next. Checkbox state is read from the
        page-independent selection model (F1); signals are blocked while
        populating so programmatic check states don't echo back into the model."""
        start = self.current_page * self.results_per_page
        page  = self.current_results[start:start + self.results_per_page]

        self.results_table.blockSignals(True)
        self.results_table.setRowCount(0)
        for row, paper in enumerate(page):
            self.results_table.insertRow(row)
            title_item = QTableWidgetItem(paper.get("title", ""))
            title_item.setData(Qt.ItemDataRole.UserRole, paper)   # store full dict for right-click
            title_item.setFlags(title_item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            title_item.setCheckState(
                Qt.CheckState.Checked
                if self._selection.is_selected(paper_key(paper))
                else Qt.CheckState.Unchecked
            )
            self.results_table.setItem(row, 0, title_item)
            authors = paper.get("authors", "")
            if isinstance(authors, list):
                authors = "; ".join(authors)
            self.results_table.setItem(row, 1, QTableWidgetItem(str(authors)[:50]))
            self.results_table.setItem(row, 2, QTableWidgetItem(paper.get("date") or paper.get("pub_date", "")))
            self.results_table.setItem(row, 3, QTableWidgetItem(paper.get("category", "")))
            self.results_table.setItem(row, 4, QTableWidgetItem(paper.get("type", paper.get("document_type", ""))))
            source_label = SOURCE_LABELS.get(paper.get("source", ""), paper.get("source", ""))
            self.results_table.setItem(row, 5, QTableWidgetItem(source_label))
        self.results_table.blockSignals(False)

        total = max(1, (len(self.current_results) + self.results_per_page - 1) // self.results_per_page)
        self.page_label.setText(f"Page {self.current_page + 1} of {total}  ({len(self.current_results):,} total)")
        self.prev_btn.setEnabled(self.current_page > 0)
        self.next_btn.setEnabled(self.current_page < total - 1)

    def prev_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.display_page()

    def next_page(self):
        total = (len(self.current_results) + self.results_per_page - 1) // self.results_per_page
        if self.current_page < total - 1:
            self.current_page += 1
            self.display_page()

    # ── Checkbox / selection helpers ──────────────────────────────────────────

    def _on_item_changed(self, item: QTableWidgetItem):
        # A user toggle of a visible row updates the page-independent model (F1).
        if item.column() == 0:
            paper = item.data(Qt.ItemDataRole.UserRole)
            if paper is not None:
                self._selection.set(
                    paper_key(paper), item.checkState() == Qt.CheckState.Checked
                )
            self._update_checked_count()

    def _update_checked_count(self):
        count = self._selection.count(self.current_results)
        self._checked_label.setText(f"{count} selected")

    def _select_all(self):
        # Select across ALL pages, not just rendered rows (F1).
        self._selection.select_all(self.current_results)
        self.display_page()  # reflect selection on the current page
        self._update_checked_count()

    def _select_none(self):
        self._selection.clear()
        self.display_page()
        self._update_checked_count()

    def _checked_papers(self) -> List[Dict[str, Any]]:
        # Return selected papers across the whole result set, not just the page.
        return self._selection.selected(self.current_results)

    def _save_selected_as_reference(self):
        papers = self._checked_papers()
        if not papers:
            QMessageBox.information(self, "Nothing selected", "Check at least one paper first.")
            return
        name, ok = QInputDialog.getText(
            self, "Save Reference List", f"Name for this list ({len(papers)} papers):"
        )
        if not ok or not name.strip():
            return
        list_id = self.db.create_reference_list(name.strip())
        if list_id is None:
            QMessageBox.warning(self, "Error", "Could not create reference list.")
            return
        added = sum(1 for p in papers if self.db.add_to_reference_list(list_id, p))
        self._select_none()
        self.reference_list_saved.emit()
        QMessageBox.information(
            self, "Saved",
            f"Saved {added} paper(s) to '{name.strip()}'.\n"
            "Open the Saved References tab to view and download."
        )

    def refresh_filters(self):
        """Called by FiltersTab after a save."""
        self.load_filters_list()


# ---------------------------------------------------------------------------
# Filters tab
# ---------------------------------------------------------------------------
# Text filter group widgets
# ---------------------------------------------------------------------------

class TextFilterGroupWidget(QWidget):
    """
    One AND-condition group. Within the group:
      - Title terms (comma-sep) — any term must appear in title (OR within field)
      - Abstract terms (comma-sep) — any term must appear in abstract
      - Title+Abstract terms (comma-sep) — any term must appear in either
    All non-empty fields must match (AND between fields).
    """
    remove_requested = pyqtSignal(object)

    def __init__(self, data: Dict[str, str] = None, parent=None):
        super().__init__(parent)
        data = data or {}
        self.setObjectName("group")

        outer = QVBoxLayout()
        outer.setContentsMargins(0, 0, 0, 0)

        # Header row: label + remove button
        header = QHBoxLayout()
        self.label = QLabel("AND group")
        self.label.setStyleSheet("font-weight: bold; color: #555;")
        remove_btn = QPushButton("✕ Remove")
        remove_btn.setFixedWidth(80)
        remove_btn.clicked.connect(lambda: self.remove_requested.emit(self))
        header.addWidget(self.label)
        header.addStretch()
        header.addWidget(remove_btn)
        outer.addLayout(header)

        form = QFormLayout()
        self.title_field    = QLineEdit(data.get("title",    ""))
        self.abstract_field = QLineEdit(data.get("abstract", ""))
        self.both_field     = QLineEdit(data.get("both",     ""))

        self.title_field.setPlaceholderText("e.g. hippocampus, memory  (comma = OR)")
        self.abstract_field.setPlaceholderText("e.g. rodent, mouse  (comma = OR)")
        self.both_field.setPlaceholderText("e.g. humans  (comma = OR)")

        form.addRow("Title contains:",           self.title_field)
        form.addRow("Abstract contains:",        self.abstract_field)
        form.addRow("Title OR Abstract contains:", self.both_field)
        outer.addLayout(form)

        # Visual separator
        line = QFrame(); line.setFrameShape(QFrame.Shape.HLine)
        line.setStyleSheet("color: #ccc;")
        outer.addWidget(line)

        self.setLayout(outer)

    def get_data(self) -> Dict[str, str]:
        return {
            "title":    self.title_field.text().strip(),
            "abstract": self.abstract_field.text().strip(),
            "both":     self.both_field.text().strip(),
        }

    def set_label(self, text: str):
        self.label.setText(text)


class TextFiltersWidget(QWidget):
    """
    Container for multiple TextFilterGroupWidgets.
    Groups are OR'd together — a paper must match at least one.
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self._groups: List[TextFilterGroupWidget] = []

        self._outer = QVBoxLayout()
        self._outer.setContentsMargins(0, 0, 0, 0)

        self._groups_layout = QVBoxLayout()
        self._groups_layout.setSpacing(4)
        self._outer.addLayout(self._groups_layout)

        add_btn = QPushButton("+ Add OR Group")
        add_btn.setFixedWidth(130)
        add_btn.clicked.connect(lambda: self.add_group())
        self._outer.addWidget(add_btn)

        self.setLayout(self._outer)

    def add_group(self, data: Dict[str, str] = None):
        widget = TextFilterGroupWidget(data, self)
        widget.remove_requested.connect(self._remove_group)
        self._groups.append(widget)
        self._groups_layout.addWidget(widget)
        self._relabel()

    def _remove_group(self, widget: TextFilterGroupWidget):
        self._groups.remove(widget)
        self._groups_layout.removeWidget(widget)
        widget.deleteLater()
        self._relabel()

    def _relabel(self):
        for i, g in enumerate(self._groups):
            label = "Match if:" if i == 0 else "— OR —"
            g.set_label(label)

    def get_groups(self) -> List[Dict[str, str]]:
        return [g.get_data() for g in self._groups]

    def load_groups(self, groups: List[Dict[str, str]]):
        # Clear existing
        for g in list(self._groups):
            self._groups_layout.removeWidget(g)
            g.deleteLater()
        self._groups.clear()
        for data in groups:
            self.add_group(data)
        if not self._groups:
            self.add_group()   # always show at least one


# ---------------------------------------------------------------------------
# Discover search terms via LLM
# ---------------------------------------------------------------------------

class DiscoverTermsWorker(QObject):
    """
    Search broadly with a natural-language description, then ask the local
    LLM to suggest precise search terms from the papers found.
    """
    status   = pyqtSignal(str)
    finished = pyqtSignal(str)    # structured text block from LLM
    error    = pyqtSignal(str)

    def __init__(self, orchestrator, query: str, days_back: int = 90, max_papers: int = 30):
        super().__init__()
        self.orchestrator = orchestrator
        self.query = query
        self.days_back = days_back
        self.max_papers = max_papers
        self._stop = False

    def stop(self):
        self._stop = True

    @staticmethod
    def _papers_fallback(papers: list) -> str:
        """Format found papers as a readable list to show when no LLM is available."""
        lines = [
            "[No LLM available — configure one in Settings > LLM Config.]\n"
            "[Papers found below. Write search terms manually in the box.]\n"
        ]
        for i, r in enumerate(papers[:30], 1):
            lines.append(f"{i}. {r.title}")
            if r.abstract:
                lines.append(f"   {r.abstract[:150].strip()}…")
        return "\n".join(lines)

    @staticmethod
    def _query_to_keywords(query: str) -> str:
        """Shared with the web app (src/discover.py); stop words from llm_config.yaml."""
        from src.discover import discover_settings, query_to_keywords
        from src.llm_config import load_llm_config
        return query_to_keywords(query, discover_settings(load_llm_config()).stop_words)

    def run(self):
        self.status.emit("Searching for relevant papers…")
        keywords = self._query_to_keywords(self.query)
        filter_dict = {
            "days_back": self.days_back,
            "text_groups": [{"title": "", "abstract": "", "both": keywords}],
        }
        papers = []

        def on_batch(records):
            papers.extend(records)

        try:
            self.orchestrator.search(
                filter_dict=filter_dict,
                on_batch=on_batch,
                should_stop=lambda: self._stop,
                max_results=self.max_papers,
            )
        except Exception as e:
            if not self._stop:
                self.error.emit(f"Search failed: {e}")
            return

        if self._stop:
            return

        if not papers:
            self.error.emit(
                "No papers found for that description. Try broader wording or a longer date range."
            )
            return

        self.status.emit(f"Found {len(papers)} papers — asking LLM for search terms…")

        paper_lines = "\n".join(
            f"- {r.title}: {(r.abstract or '')[:200].strip()}"
            for r in papers[:25]
        )
        prompt = (
            f"Research interest: {self.query}\n\n"
            f"Sample papers:\n{paper_lines}\n\n"
            "Based on these papers, suggest search terms for finding similar research. "
            "Organise them into 4-6 thematic groups. Use this exact format:\n\n"
            "**Group name**\n"
            "- term one / synonym / abbreviation\n"
            "- term two\n"
            "- term three\n\n"
            "Rules:\n"
            "- Group names must be wrapped in **double asterisks**\n"
            "- Each term line must start with '- '\n"
            "- Use ' / ' to separate close synonyms or abbreviations on the same line\n"
            "- Focus on methodology, technical vocabulary, domain-specific constructs\n"
            "- Avoid generic words\n"
            "- No preamble, no explanation — output only the groups and terms"
        )

        from src.llm_providers import resolve_client, NoLLMCredentialError
        _settings_hint = (
            "To configure: open the Settings tab → select 'LLM Config' → "
            "set default_provider to 'anthropic' or 'deepseek', and add your api_key."
        )
        try:
            resolved = resolve_client()
        except NoLLMCredentialError as e:
            logger.warning("Discover Terms: no LLM configured (%s)", e)
            self.status.emit(f"No LLM key set. {_settings_hint}")
            self.finished.emit(self._papers_fallback(papers))
            return

        if not resolved.client.is_available():
            logger.warning("Discover Terms: provider %r not reachable", resolved.provider)
            self.status.emit(
                f"LLM provider '{resolved.provider}' is not running. "
                f"Start Ollama, or {_settings_hint}"
            )
            self.finished.emit(self._papers_fallback(papers))
            return

        result = resolved.client.generate(prompt)
        if not result:
            logger.warning("Discover Terms: LLM (%s) returned empty response", resolved.provider)
            self.status.emit(f"LLM ({resolved.provider}) returned no response — showing papers.")
            self.finished.emit(self._papers_fallback(papers))
            return

        self.finished.emit(result.strip())


class DiscoverTermsDialog(QDialog):
    """Let the user describe their research interest and receive suggested search terms."""

    def __init__(self, orchestrator, parent=None):
        super().__init__(parent)
        self.orchestrator = orchestrator
        self.setWindowTitle("Discover Search Terms")
        self.setMinimumSize(560, 480)
        self._thread = None
        self._worker = None
        self._init_ui()

    def _init_ui(self):
        layout = QVBoxLayout()

        layout.addWidget(QLabel(
            "Describe your research interest in plain language. The app will search "
            "for relevant papers and ask the local LLM to suggest search terms organised "
            "by theme."
        ))

        self.query_edit = QTextEdit()
        self.query_edit.setPlaceholderText(
            "e.g. agents modeling social behavior or social interactions"
        )
        self.query_edit.setMaximumHeight(72)
        layout.addWidget(self.query_edit)

        days_row = QHBoxLayout()
        days_row.addWidget(QLabel("Look back (days):"))
        self.days_spin = QSpinBox()
        self.days_spin.setRange(7, 3650)
        self.days_spin.setValue(180)
        days_row.addWidget(self.days_spin)
        days_row.addStretch()
        self.discover_btn = QPushButton("Discover")
        self.discover_btn.setDefault(True)
        self.discover_btn.clicked.connect(self._start)
        days_row.addWidget(self.discover_btn)
        layout.addLayout(days_row)

        self.status_label = QLabel("")
        self.status_label.setStyleSheet("color: gray; font-size: 11px;")
        layout.addWidget(self.status_label)

        # ── Structured output (read-only) ────────────────────────────────────
        layout.addWidget(QLabel("Suggested terms by theme:"))
        self.structured_view = QTextEdit()
        self.structured_view.setReadOnly(True)
        self.structured_view.setPlaceholderText(
            "Themes and terms will appear here after discovery…"
        )
        self.structured_view.setFontFamily("Courier")
        layout.addWidget(self.structured_view, 3)

        # ── Editable flat list for the filter ───────────────────────────────
        filter_hdr = QHBoxLayout()
        filter_hdr.addWidget(QLabel("Terms for filter (comma-separated, edit as needed):"))
        self.parse_btn = QPushButton("Re-parse from above")
        self.parse_btn.setEnabled(False)
        self.parse_btn.setToolTip("Re-extract terms from the theme list above")
        self.parse_btn.clicked.connect(self._populate_filter_terms)
        filter_hdr.addStretch()
        filter_hdr.addWidget(self.parse_btn)
        layout.addLayout(filter_hdr)

        self.terms_edit = QTextEdit()
        self.terms_edit.setPlaceholderText("Comma-separated terms will appear here…")
        self.terms_edit.setMaximumHeight(72)
        layout.addWidget(self.terms_edit)

        btns = QHBoxLayout()
        self.apply_btn = QPushButton("Add to filter as new group")
        self.apply_btn.setEnabled(False)
        self.apply_btn.clicked.connect(self.accept)
        close_btn = QPushButton("Close")
        close_btn.clicked.connect(self.reject)
        btns.addStretch()
        btns.addWidget(close_btn)
        btns.addWidget(self.apply_btn)
        layout.addLayout(btns)

        self.setLayout(layout)

    def _start(self):
        query = self.query_edit.toPlainText().strip()
        if not query:
            self.status_label.setText("Enter a research description first.")
            return

        self._stop_worker()
        self.discover_btn.setEnabled(False)
        self.apply_btn.setEnabled(False)
        self.parse_btn.setEnabled(False)
        self.structured_view.setPlainText("")
        self.terms_edit.setPlainText("")
        self.status_label.setText("Starting…")

        self._worker = DiscoverTermsWorker(
            self.orchestrator, query, days_back=self.days_spin.value()
        )
        self._thread = QThread()
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.status.connect(self.status_label.setText)
        self._worker.finished.connect(self._on_finished)
        self._worker.error.connect(self._on_error)
        self._thread.start()

    def _on_finished(self, structured_text: str):
        self._thread.quit()
        self.discover_btn.setEnabled(True)
        if not structured_text:
            self.status_label.setText("Done (no results).")
            return
        self.structured_view.setPlainText(structured_text)
        # Only auto-parse and enable "Add" when the LLM returned themed output
        # (identified by **headers**).  The fallback paper list is shown for
        # reference only; the user types terms manually in the box below.
        is_llm_output = "**" in structured_text
        if is_llm_output:
            self._populate_filter_terms()
            self.parse_btn.setEnabled(True)
            self.apply_btn.setEnabled(True)
            self.status_label.setText("Done — edit the filter terms below, then click Add.")
        else:
            self.terms_edit.setPlaceholderText(
                "LLM not available — type search terms here (comma-separated)…"
            )
            self.apply_btn.setEnabled(True)   # let them add manual terms
            self.status_label.setText(
                "LLM not available. Browse the papers above, type terms below, then click Add."
            )

    @staticmethod
    def _parse_structured(text: str) -> str:
        """
        Extract individual search terms from a structured LLM block.

        Expected format (requested in the prompt):
            **Group name**
            - term one / synonym / abbreviation
            - term two

        Lines starting with '**' are group headers — skipped.
        Lines starting with '-' (after stripping) are term lines.
        Slash-separated variants on the same line become separate terms.
        Parenthetical abbreviations like "(ABSS)" are stripped from each variant.
        Lines matching neither pattern are skipped.
        """
        import re
        terms = []
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            # Skip group headers (**...**)
            if line.startswith("**"):
                continue
            # Accept term lines (start with '-' or '•')
            if line.startswith(("-", "•")):
                line = line.lstrip("-•").strip()
            else:
                # Tolerate lines without a leading dash if they contain a slash
                # (some models omit the bullet); skip everything else
                if "/" not in line:
                    continue
            # Split slash variants
            parts = re.split(r"\s*/\s*", line)
            for p in parts:
                p = re.sub(r"\s*\([^)]+\)", "", p).strip().strip('"').strip("'")
                if p and len(p) > 2:
                    terms.append(p)
        # Deduplicate preserving order
        seen: set = set()
        unique = []
        for t in terms:
            key = t.lower()
            if key not in seen:
                seen.add(key)
                unique.append(t)
        return ", ".join(unique)

    def _populate_filter_terms(self):
        structured = self.structured_view.toPlainText().strip()
        if structured:
            self.terms_edit.setPlainText(self._parse_structured(structured))

    def _on_error(self, msg):
        self._thread.quit()
        self.discover_btn.setEnabled(True)
        self.status_label.setText(f"Error: {msg}")

    def _stop_worker(self):
        if self._worker:
            self._worker.stop()
        if self._thread and self._thread.isRunning():
            self._thread.quit()
            self._thread.wait(3000)

    def closeEvent(self, event):
        self._stop_worker()
        super().closeEvent(event)

    def get_terms_text(self) -> str:
        return self.terms_edit.toPlainText().strip()


# ---------------------------------------------------------------------------

class FiltersTab(QWidget):
    """Build, save, test, and manage search filters."""

    filters_changed = pyqtSignal()   # notify Search tab to reload

    def __init__(self, orchestrator: "SourceOrchestrator"):
        super().__init__()
        self.orchestrator = orchestrator
        self.filters: List[Dict[str, Any]] = load_filters(FILTERS_PATH)
        self._current_index = -1
        self._threads: list = []
        # Source picker for filter forms
        enabled  = orchestrator.get_enabled_sources()
        defaults = enabled  # all enabled by default
        self._source_picker = SourcePickerWidget(enabled, defaults)
        self.init_ui()
        self._populate_list()

    # ── UI build ─────────────────────────────────────────────────────────────

    def init_ui(self):
        outer = QHBoxLayout()

        # ── Left: filter list ────────────────────────────────────────────────
        left = QWidget(); left.setMaximumWidth(220)
        ll   = QVBoxLayout()
        ll.addWidget(QLabel("Saved Filters:"))

        self.filter_list = QListWidget()
        self.filter_list.currentRowChanged.connect(self._on_select)
        ll.addWidget(self.filter_list)

        left.setLayout(ll)

        # ── Right: form + test results ───────────────────────────────────────
        right = QWidget()
        rl    = QVBoxLayout()

        # ─ Discover search terms (top of page) ─
        discover_group = QGroupBox("Discover Search Terms with AI")
        dg_layout = QHBoxLayout()
        dg_label = QLabel(
            "Describe your research interest — the app will search for relevant papers "
            "and ask the LLM to suggest keywords."
        )
        dg_label.setWordWrap(True)
        dg_btn = QPushButton("🔍  Discover terms…")
        dg_btn.setToolTip("Configure LLM in Settings > LLM Config")
        dg_btn.setFixedWidth(160)
        dg_btn.clicked.connect(self._open_discover_dialog)
        dg_layout.addWidget(dg_label, 1)
        dg_layout.addWidget(dg_btn)
        discover_group.setLayout(dg_layout)
        rl.addWidget(discover_group)

        # ─ Identity ─
        id_group = QGroupBox("Filter Identity")
        id_form  = QFormLayout()
        self.name_field    = QLineEdit()
        self.enabled_check = QCheckBox("Enabled (will run with 'Run All Enabled')")
        id_form.addRow("Name:", self.name_field)
        id_form.addRow("",      self.enabled_check)
        id_group.setLayout(id_form)
        rl.addWidget(id_group)

        # ─ Scope ─
        scope_group = QGroupBox("Scope")
        scope_form  = QFormLayout()
        self.category_combo = QComboBox(); self.category_combo.addItems(BIORXIV_CATEGORIES)

        # Days-back spinner (default mode)
        self.days_spin = QSpinBox(); self.days_spin.setRange(1, 3650); self.days_spin.setValue(7)

        # Date-range pickers (hidden until toggled on)
        today = QDate.currentDate()
        self.start_date_edit = QDateEdit(today.addDays(-7)); self.start_date_edit.setCalendarPopup(True)
        self.end_date_edit   = QDateEdit(today);             self.end_date_edit.setCalendarPopup(True)
        self.start_date_edit.setDisplayFormat("yyyy-MM-dd")
        self.end_date_edit.setDisplayFormat("yyyy-MM-dd")

        # Toggle checkbox
        self.date_range_check = QCheckBox("Use date range")
        self.date_range_check.toggled.connect(self._on_date_range_toggled)

        # Container so we can swap rows cleanly
        self._days_widget  = self.days_spin
        self._range_widget = QWidget()
        range_row = QHBoxLayout(); range_row.setContentsMargins(0, 0, 0, 0)
        range_row.addWidget(QLabel("From:")); range_row.addWidget(self.start_date_edit)
        range_row.addWidget(QLabel("To:"));   range_row.addWidget(self.end_date_edit)
        self._range_widget.setLayout(range_row)
        self._range_widget.setVisible(False)

        scope_form.addRow("Category:",    self.category_combo)
        scope_form.addRow("",             self.date_range_check)
        scope_form.addRow("Days back:",   self._days_widget)
        scope_form.addRow("Date range:",  self._range_widget)
        scope_group.setLayout(scope_form)
        rl.addWidget(scope_group)

        # ─ Sources ─
        sources_group = QGroupBox("Sources")
        sg_layout = QHBoxLayout()
        sg_layout.addWidget(self._source_picker)
        sources_group.setLayout(sg_layout)
        rl.addWidget(sources_group)

        # ─ Text filters ─
        text_group = QGroupBox(
            "Text Search Groups  —  OR between groups, AND within a group  —  comma = OR within a field"
        )
        tg_layout = QVBoxLayout()
        self.text_filters = TextFiltersWidget()
        tg_layout.addWidget(self.text_filters)
        text_group.setLayout(tg_layout)
        rl.addWidget(text_group)

        # ─ Other text filters ─
        other_group = QGroupBox("Author / Institution  (comma-separated — any match wins)")
        other_form  = QFormLayout()
        self.authors_field     = QLineEdit(); self.authors_field.setPlaceholderText("e.g. Smith, Doudna")
        self.institution_field = QLineEdit(); self.institution_field.setPlaceholderText("e.g. Broad Institute")
        other_form.addRow("Authors:",     self.authors_field)
        other_form.addRow("Institution:", self.institution_field)
        other_group.setLayout(other_form)
        rl.addWidget(other_group)

        # ─ Structured filters ─
        struct_group = QGroupBox("Structured Filters")
        struct_form  = QFormLayout()
        self.type_combo      = QComboBox(); self.type_combo.addItems(PAPER_TYPES)
        self.version_combo   = QComboBox(); self.version_combo.addItems(VERSIONS)
        self.published_combo = QComboBox(); self.published_combo.addItems(PUBLISHED)
        self.license_combo   = QComboBox(); self.license_combo.addItems(LICENSES)
        self.species_combo   = QComboBox(); self.species_combo.addItems(SPECIES)
        struct_form.addRow("Paper type:", self.type_combo)
        struct_form.addRow("Version:",    self.version_combo)
        struct_form.addRow("Published:",  self.published_combo)
        struct_form.addRow("License:",    self.license_combo)
        struct_form.addRow("Study type:", self.species_combo)
        struct_group.setLayout(struct_form)
        rl.addWidget(struct_group)

        # ─ Action buttons ─
        action_row = QHBoxLayout()
        new_btn     = QPushButton("+ New");          new_btn.clicked.connect(self._new_filter)
        save_btn    = QPushButton("💾  Save");       save_btn.clicked.connect(self._save_filter)
        save_as_btn = QPushButton("📋  Save As");    save_as_btn.clicked.connect(self._save_as_filter)
        del_btn     = QPushButton("🗑  Delete");     del_btn.clicked.connect(self._delete_filter)
        self.test_btn = QPushButton("▶  Test Filter"); self.test_btn.clicked.connect(self._test_filter)
        action_row.addWidget(new_btn)
        action_row.addWidget(save_btn)
        action_row.addWidget(save_as_btn)
        action_row.addWidget(del_btn)
        action_row.addStretch()
        action_row.addWidget(self.test_btn)
        rl.addLayout(action_row)

        prog_row = QHBoxLayout()
        self.test_status = QLabel("")
        self.test_progress = QProgressBar()
        self.test_progress.setTextVisible(True)
        self.test_progress.setFormat("Fetched %v / %m papers")
        self.test_progress.setVisible(False)
        prog_row.addWidget(self.test_status)
        prog_row.addWidget(self.test_progress)
        rl.addLayout(prog_row)

        # ─ Test results table ─
        rl.addWidget(QLabel("Test Results:"))
        self.test_table = QTableWidget()
        self.test_table.setColumnCount(6)
        self.test_table.setHorizontalHeaderLabels(["Title", "Authors", "Date", "Category", "Type", "Source"])
        self.test_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.test_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.test_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.test_table.setWordWrap(True)
        self.test_table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.test_table.setToolTip("Right-click a row to view abstract/discussion or download PDF")
        _attach_context_menu(self.test_table)
        rl.addWidget(self.test_table)

        right.setLayout(rl)

        scroll = QScrollArea()
        scroll.setWidget(right)
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(left)
        splitter.addWidget(scroll)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        outer.addWidget(splitter)
        self.setLayout(outer)

    # ── Date-range toggle ─────────────────────────────────────────────────────

    def _on_date_range_toggled(self, checked: bool):
        self._days_widget.setVisible(not checked)
        self._range_widget.setVisible(checked)

    # ── List management ───────────────────────────────────────────────────────

    def _populate_list(self):
        self.filter_list.clear()
        for f in self.filters:
            self.filter_list.addItem(f.get("name", "Unnamed"))
        if self.filters:
            self.filter_list.setCurrentRow(0)

    def _on_select(self, index: int):
        if index < 0 or index >= len(self.filters):
            return
        self._current_index = index
        self._load_form(self.filters[index])

    def _load_form(self, f: Dict[str, Any]):
        self.name_field.setText(f.get("name", ""))
        self.enabled_check.setChecked(f.get("enabled", True))
        cat = f.get("category", "(any)")
        self.category_combo.setCurrentText(cat if cat in BIORXIV_CATEGORIES else "(any)")
        # Date scope
        start = f.get("start_date", "")
        end   = f.get("end_date", "")
        use_range = bool(start and end)
        self.date_range_check.setChecked(use_range)
        if use_range:
            self.start_date_edit.setDate(QDate.fromString(start, "yyyy-MM-dd"))
            self.end_date_edit.setDate(QDate.fromString(end,   "yyyy-MM-dd"))
        else:
            self.days_spin.setValue(f.get("days_back", 7))
        # Load text groups — back-compat: migrate old flat keywords list
        groups = f.get("text_groups", [])
        if not groups and f.get("keywords"):
            groups = [{"title": "", "abstract": "", "both": ", ".join(f["keywords"])}]
        self.text_filters.load_groups(groups)
        self.authors_field.setText(", ".join(f.get("authors", [])))
        self.institution_field.setText(f.get("institution", ""))
        self.type_combo.setCurrentText(f.get("paper_type", "(any)"))
        self.version_combo.setCurrentText(f.get("version", "(any)"))
        self.published_combo.setCurrentText(f.get("published", "(any)"))
        self.license_combo.setCurrentText(f.get("license", "(any)"))
        self.species_combo.setCurrentText(f.get("species", "(any)"))
        self._source_picker.load_selection(
            f.get("source_selection", {"all": True, "selected": []})
        )
        self.test_table.setRowCount(0)
        self.test_status.setText("")

    def _form_to_dict(self) -> Dict[str, Any]:
        use_range = self.date_range_check.isChecked()
        return {
            "name":             self.name_field.text().strip() or "Unnamed",
            "enabled":          self.enabled_check.isChecked(),
            "category":         self.category_combo.currentText(),
            "days_back":        0 if use_range else self.days_spin.value(),
            "start_date":       self.start_date_edit.date().toString("yyyy-MM-dd") if use_range else "",
            "end_date":         self.end_date_edit.date().toString("yyyy-MM-dd")   if use_range else "",
            "text_groups":      self.text_filters.get_groups(),
            "authors":          [a.strip() for a in self.authors_field.text().split(",") if a.strip()],
            "institution":      self.institution_field.text().strip(),
            "paper_type":       self.type_combo.currentText(),
            "version":          self.version_combo.currentText(),
            "published":        self.published_combo.currentText(),
            "license":          self.license_combo.currentText(),
            "species":          self.species_combo.currentText(),
            "source_selection": self._source_picker.get_selection(),
        }

    # ── CRUD ──────────────────────────────────────────────────────────────────

    def _new_filter(self):
        blank = {
            "name": "New Filter", "enabled": True, "category": "(any)",
            "days_back": 7, "start_date": "", "end_date": "",
            "text_groups": [], "authors": [], "institution": "",
            "paper_type": "(any)", "version": "(any)", "published": "(any)", "license": "(any)",
            "species": "(any)", "source_selection": {"all": True, "selected": []},
        }
        self.filters.append(blank)
        save_filters(FILTERS_PATH, self.filters)
        self._populate_list()
        self.filter_list.setCurrentRow(len(self.filters) - 1)

    def _save_filter(self):
        f = self._form_to_dict()
        if not _filter_has_text(f):
            self.test_status.setStyleSheet("font-weight: bold; color: #cc4400;")
            self.test_status.setText("⚠  Cannot save: add at least one search term first.")
            return
        self.test_status.setStyleSheet("")
        # If the name was changed, treat as a new filter rather than overwriting
        original_name = (
            self.filters[self._current_index].get("name", "")
            if self._current_index >= 0 else ""
        )
        name_changed = self._current_index >= 0 and f["name"] != original_name

        if self._current_index < 0 or name_changed:
            self.filters.append(f)
            self._current_index = len(self.filters) - 1
            save_filters(FILTERS_PATH, self.filters)
            self._populate_list()
            self.filter_list.setCurrentRow(self._current_index)
            self.test_status.setText(f"Saved as new filter '{f['name']}'.")
        else:
            self.filters[self._current_index] = f
            save_filters(FILTERS_PATH, self.filters)
            self.test_status.setText("Filter saved.")
        self.filters_changed.emit()

    def _save_as_filter(self):
        """Duplicate current form as a brand-new filter with a prompted name."""
        f = self._form_to_dict()
        if not _filter_has_text(f):
            self.test_status.setStyleSheet("font-weight: bold; color: #cc4400;")
            self.test_status.setText("⚠  Cannot save: add at least one search term first.")
            return
        self.test_status.setStyleSheet("")
        new_name, ok = QInputDialog.getText(
            self, "Save As", "New filter name:", text=f["name"] + " (copy)"
        )
        if not ok or not new_name.strip():
            return
        f["name"] = new_name.strip()
        self.filters.append(f)
        save_filters(FILTERS_PATH, self.filters)
        self._populate_list()
        self._current_index = len(self.filters) - 1
        self.filter_list.setCurrentRow(self._current_index)
        self.test_status.setText(f"Saved as '{f['name']}'.")
        self.filters_changed.emit()

    def _delete_filter(self):
        if self._current_index < 0 or not self.filters:
            return
        name = self.filters[self._current_index].get("name", "this filter")
        if QMessageBox.question(self, "Delete", f"Delete '{name}'?") != QMessageBox.StandardButton.Yes:
            return
        del self.filters[self._current_index]
        save_filters(FILTERS_PATH, self.filters)
        self._current_index = -1
        self._populate_list()
        self.test_table.setRowCount(0)
        self.test_status.setText("")
        self.filters_changed.emit()

    # ── Test ──────────────────────────────────────────────────────────────────

    def _test_filter(self):
        f = self._form_to_dict()
        self.test_table.setRowCount(0)
        if not _filter_has_text(f):
            self.test_status.setStyleSheet("font-weight: bold; color: #cc4400;")
            self.test_status.setText(
                "⚠  No search terms configured. "
                "Add keywords in the Text Search Groups before testing."
            )
            return
        self.test_progress.setMaximum(0)          # indeterminate / pulsing
        self.test_progress.setFormat("Connecting…")
        self.test_progress.setVisible(True)
        self.test_status.setStyleSheet("font-weight: bold; color: #0055aa;")
        self.test_status.setText("⏳  Searching…")
        self.test_btn.setText("⏹  Stop")
        self.test_btn.clicked.disconnect()
        self.test_btn.clicked.connect(self._stop_test)

        self._test_worker = SearchWorker(self.orchestrator, f, save_to_db=False)
        thread = QThread()
        self._test_worker.moveToThread(thread)
        thread.started.connect(self._test_worker.run)
        self._test_worker.batch_ready.connect(self._append_test_batch)
        self._test_worker.progress.connect(self._update_test_progress)
        self._test_worker.status.connect(self.test_status.setText)
        self._test_worker.finished.connect(lambda papers: self._on_test_done(papers))
        self._test_worker.error.connect(lambda e: self._on_test_done([]))
        thread.start()
        self._threads.append((thread, self._test_worker))

    def _stop_test(self):
        if hasattr(self, "_test_worker") and self._test_worker:
            self._test_worker.stop()
        self.test_btn.setEnabled(False)

    def _update_test_progress(self, fetched: int, total: int):
        if total > 0:
            self.test_progress.setMaximum(total)
            self.test_progress.setFormat("Fetched %v / %m papers")
            self.test_progress.setValue(fetched)

    def _append_test_batch(self, papers: list):
        for p in papers:
            row = self.test_table.rowCount()
            self.test_table.insertRow(row)
            title_item = QTableWidgetItem(p.get("title", ""))
            title_item.setData(Qt.ItemDataRole.UserRole, p)
            self.test_table.setItem(row, 0, title_item)
            authors = p.get("authors", "")
            if isinstance(authors, list):
                authors = "; ".join(authors)
            self.test_table.setItem(row, 1, QTableWidgetItem(str(authors)[:50]))
            self.test_table.setItem(row, 2, QTableWidgetItem(p.get("date") or p.get("pub_date", "")))
            self.test_table.setItem(row, 3, QTableWidgetItem(p.get("category", "")))
            self.test_table.setItem(row, 4, QTableWidgetItem(p.get("type", p.get("document_type", ""))))
            source_label = SOURCE_LABELS.get(p.get("source", ""), p.get("source", ""))
            self.test_table.setItem(row, 5, QTableWidgetItem(source_label))

    def _on_test_done(self, papers: list):
        self.test_progress.setVisible(False)
        self.test_status.setStyleSheet("")
        self.test_status.setText(f"{len(papers):,} papers matched")
        self.test_btn.setText("▶  Test Filter")
        self.test_btn.setEnabled(True)
        self.test_btn.clicked.disconnect()
        self.test_btn.clicked.connect(self._test_filter)

    def _show_test_results(self, papers: list):
        # Legacy - no longer called directly but kept for safety
        self.test_table.setRowCount(0)
        for row, p in enumerate(papers[:50]):
            self.test_table.insertRow(row)
            title_item = QTableWidgetItem(p.get("title", ""))
            title_item.setData(Qt.ItemDataRole.UserRole, p)   # store full dict for right-click
            self.test_table.setItem(row, 0, title_item)
            authors = p.get("authors", "")
            if isinstance(authors, list):
                authors = "; ".join(authors)
            self.test_table.setItem(row, 1, QTableWidgetItem(str(authors)[:50]))
            self.test_table.setItem(row, 2, QTableWidgetItem(p.get("date") or p.get("pub_date", "")))
            self.test_table.setItem(row, 3, QTableWidgetItem(p.get("category", "")))
            self.test_table.setItem(row, 4, QTableWidgetItem(p.get("type", "")))

    def _open_discover_dialog(self):
        dlg = DiscoverTermsDialog(self.orchestrator, parent=self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            terms_text = dlg.get_terms_text()
            if terms_text:
                self.text_filters.add_group({"both": terms_text})


# ---------------------------------------------------------------------------
# Main window
# ---------------------------------------------------------------------------
# Saved References tab
# ---------------------------------------------------------------------------

class SavedReferencesTab(QWidget):
    """Browse, manage, and batch-download saved reference lists."""

    def __init__(self, db: Database):
        super().__init__()
        self.db = db
        self._current_list_id: Optional[int] = None
        self._download_thread: Optional[QThread] = None
        self._download_worker: Optional[BatchPdfDownloadWorker] = None
        self.init_ui()
        self.refresh_lists()

    def init_ui(self):
        outer = QHBoxLayout()

        # ── Left: list of saved collections ─────────────────────────────────
        left = QWidget(); left.setMaximumWidth(240)
        ll   = QVBoxLayout()
        ll.addWidget(QLabel("Reference Lists:"))
        self.lists_widget = QListWidget()
        self.lists_widget.currentRowChanged.connect(self._on_list_selected)
        ll.addWidget(self.lists_widget)
        del_list_btn = QPushButton("🗑  Delete List")
        del_list_btn.clicked.connect(self._delete_list)
        ll.addWidget(del_list_btn)
        left.setLayout(ll)

        # ── Right: papers in selected list ───────────────────────────────────
        right = QWidget()
        rl    = QVBoxLayout()

        self.list_title_label = QLabel("Select a list to view its papers")
        self.list_title_label.setStyleSheet("font-weight: bold; font-size: 13px;")
        rl.addWidget(self.list_title_label)

        self.papers_table = QTableWidget()
        self.papers_table.setColumnCount(5)
        self.papers_table.setHorizontalHeaderLabels(["Title", "Authors", "Date", "Type", "Source"])
        self.papers_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.papers_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.papers_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.papers_table.setWordWrap(True)
        self.papers_table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.papers_table.itemChanged.connect(self._on_paper_item_changed)
        _attach_context_menu(self.papers_table)
        rl.addWidget(self.papers_table, 5)

        # Selection bar
        sel_bar = QHBoxLayout()
        sel_all_btn  = QPushButton("☑ Select All");  sel_all_btn.clicked.connect(self._select_all)
        sel_none_btn = QPushButton("☐ Clear");        sel_none_btn.clicked.connect(self._select_none)
        self._ref_checked_label = QLabel("0 selected")
        sel_bar.addWidget(sel_all_btn)
        sel_bar.addWidget(sel_none_btn)
        sel_bar.addWidget(self._ref_checked_label)
        sel_bar.addStretch()
        remove_btn = QPushButton("Remove Selected from List")
        remove_btn.clicked.connect(self._remove_selected)
        sel_bar.addWidget(remove_btn)
        rl.addLayout(sel_bar)

        # Download bar
        dl_bar = QHBoxLayout()
        dl_sel_btn = QPushButton("⬇  Download Selected PDFs")
        dl_sel_btn.clicked.connect(lambda: self._start_download(selected_only=True))
        dl_all_btn = QPushButton("⬇  Download All PDFs")
        dl_all_btn.clicked.connect(lambda: self._start_download(selected_only=False))
        export_btn = QPushButton("📑  Export to Excel…")
        export_btn.setToolTip("Export this list (or selected papers) to an .xlsx file with clickable links")
        export_btn.clicked.connect(self._export_to_excel)
        self._stop_dl_btn = QPushButton("⏹  Stop")
        self._stop_dl_btn.setVisible(False)
        self._stop_dl_btn.clicked.connect(self._stop_download)
        dl_bar.addWidget(dl_sel_btn)
        dl_bar.addWidget(dl_all_btn)
        dl_bar.addWidget(export_btn)
        dl_bar.addStretch()
        dl_bar.addWidget(self._stop_dl_btn)
        rl.addLayout(dl_bar)

        # Progress
        self._dl_status = QLabel("")
        self._dl_progress = QProgressBar()
        self._dl_progress.setTextVisible(True)
        self._dl_progress.setVisible(False)
        rl.addWidget(self._dl_status)
        rl.addWidget(self._dl_progress)

        right.setLayout(rl)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(left)
        splitter.addWidget(right)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        outer.addWidget(splitter)
        self.setLayout(outer)

    # ── List management ───────────────────────────────────────────────────────

    def refresh_lists(self):
        self.lists_widget.clear()
        for lst in self.db.get_reference_lists():
            item = QListWidgetItem(f"{lst['name']}  ({lst['item_count']})")
            item.setData(Qt.ItemDataRole.UserRole, lst["id"])
            self.lists_widget.addItem(item)

    def _on_list_selected(self, index: int):
        if index < 0:
            self._current_list_id = None
            self.papers_table.setRowCount(0)
            self.list_title_label.setText("Select a list to view its papers")
            return
        item = self.lists_widget.item(index)
        self._current_list_id = item.data(Qt.ItemDataRole.UserRole)
        name = item.text()
        self.list_title_label.setText(name)
        self._load_papers(self._current_list_id)

    def _load_papers(self, list_id: int):
        self.papers_table.itemChanged.disconnect(self._on_paper_item_changed)
        self.papers_table.setRowCount(0)
        for entry in self.db.get_reference_list_items(list_id):
            paper = entry["paper"]
            row   = self.papers_table.rowCount()
            self.papers_table.insertRow(row)
            title_item = QTableWidgetItem(paper.get("title", ""))
            title_item.setData(Qt.ItemDataRole.UserRole, {"paper": paper, "item_id": entry["id"]})
            title_item.setFlags(title_item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            title_item.setCheckState(Qt.CheckState.Unchecked)
            self.papers_table.setItem(row, 0, title_item)
            authors = paper.get("authors", "")
            if isinstance(authors, list):
                authors = "; ".join(authors)
            self.papers_table.setItem(row, 1, QTableWidgetItem(str(authors)[:50]))
            self.papers_table.setItem(row, 2, QTableWidgetItem(paper.get("date") or paper.get("pub_date", "")))
            self.papers_table.setItem(row, 3, QTableWidgetItem(paper.get("type", paper.get("document_type", ""))))
            source_label = SOURCE_LABELS.get(paper.get("source", ""), paper.get("source", ""))
            self.papers_table.setItem(row, 4, QTableWidgetItem(source_label))
        self.papers_table.itemChanged.connect(self._on_paper_item_changed)
        self._update_checked_count()

    def _delete_list(self):
        if self._current_list_id is None:
            return
        row = self.lists_widget.currentRow()
        name = self.lists_widget.item(row).text() if row >= 0 else "this list"
        if QMessageBox.question(self, "Delete", f"Delete '{name}'?") != QMessageBox.StandardButton.Yes:
            return
        self.db.delete_reference_list(self._current_list_id)
        self._current_list_id = None
        self.refresh_lists()
        self.papers_table.setRowCount(0)
        self.list_title_label.setText("Select a list to view its papers")

    # ── Checkbox helpers ──────────────────────────────────────────────────────

    def _on_paper_item_changed(self, item: QTableWidgetItem):
        if item.column() == 0:
            self._update_checked_count()

    def _update_checked_count(self):
        count = sum(
            1 for r in range(self.papers_table.rowCount())
            if (it := self.papers_table.item(r, 0)) and it.checkState() == Qt.CheckState.Checked
        )
        self._ref_checked_label.setText(f"{count} selected")

    def _select_all(self):
        self.papers_table.itemChanged.disconnect(self._on_paper_item_changed)
        for r in range(self.papers_table.rowCount()):
            it = self.papers_table.item(r, 0)
            if it:
                it.setCheckState(Qt.CheckState.Checked)
        self.papers_table.itemChanged.connect(self._on_paper_item_changed)
        self._update_checked_count()

    def _select_none(self):
        self.papers_table.itemChanged.disconnect(self._on_paper_item_changed)
        for r in range(self.papers_table.rowCount()):
            it = self.papers_table.item(r, 0)
            if it:
                it.setCheckState(Qt.CheckState.Unchecked)
        self.papers_table.itemChanged.connect(self._on_paper_item_changed)
        self._update_checked_count()

    def _checked_entries(self):
        entries = []
        for r in range(self.papers_table.rowCount()):
            it = self.papers_table.item(r, 0)
            if it and it.checkState() == Qt.CheckState.Checked:
                entries.append(it.data(Qt.ItemDataRole.UserRole))
        return entries

    def _remove_selected(self):
        entries = self._checked_entries()
        if not entries:
            return
        for e in entries:
            self.db.remove_from_reference_list(e["item_id"])
        if self._current_list_id:
            self._load_papers(self._current_list_id)
            self.refresh_lists()

    # ── Export ──────────────────────────────────────────────────────────────

    def _export_to_excel(self):
        """Export the current list (selected papers if any are checked, else all)
        to an .xlsx file with a clickable link per paper."""
        if self._current_list_id is None:
            QMessageBox.information(self, "No list", "Select a reference list first.")
            return

        checked = self._checked_entries()
        papers = (
            [e["paper"] for e in checked]
            if checked
            else self._papers_for_download(selected_only=False)
        )
        if not papers:
            QMessageBox.information(self, "Nothing to export", "This list has no papers.")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError:
            QMessageBox.critical(
                self, "Missing dependency",
                "Excel export needs the 'openpyxl' package.\n\n"
                "Install it with:\n    pip install openpyxl",
            )
            return

        from PyQt6.QtWidgets import QFileDialog

        # Default filename from the list name (strip the trailing " (count)").
        row = self.lists_widget.currentRow()
        raw_name = self.lists_widget.item(row).text() if row >= 0 else "references"
        base = re.sub(r"\s*\(\d+\)\s*$", "", raw_name).strip() or "references"
        safe = re.sub(r"[^\w\-. ]+", "_", base)
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Reference List", f"{safe}.xlsx", "Excel Workbook (*.xlsx)"
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        headers = ["Title", "Authors", "Date", "Type", "Source", "DOI", "Link"]
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = base[:31] or "References"  # Excel caps sheet names at 31 chars

            ws.append(headers)
            for c in range(1, len(headers) + 1):
                ws.cell(row=1, column=c).font = Font(bold=True)

            link_font = Font(color="0563C1", underline="single")
            for paper in papers:
                authors = paper.get("authors", "")
                if isinstance(authors, list):
                    authors = "; ".join(str(a) for a in authors)
                source = SOURCE_LABELS.get(paper.get("source", ""), paper.get("source", ""))
                ws.append([
                    paper.get("title", ""),
                    str(authors),
                    paper.get("date") or paper.get("pub_date", ""),
                    paper.get("type", paper.get("document_type", "")),
                    source,
                    paper.get("doi", ""),
                    "",  # link cell filled below as a hyperlink
                ])
                cell = ws.cell(row=ws.max_row, column=len(headers))
                url = _paper_link(paper)
                if url:
                    cell.value = "Open paper"
                    cell.hyperlink = url
                    cell.font = link_font

            # Reasonable column widths.
            widths = [60, 30, 12, 12, 16, 28, 14]
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"

            wb.save(path)
        except Exception as e:
            logger.error("Excel export failed: %s", e)
            QMessageBox.critical(self, "Export failed", f"Could not write the file:\n{e}")
            return

        QMessageBox.information(
            self, "Export complete",
            f"Exported {len(papers)} paper(s) to:\n{path}",
        )

    # ── Batch download ────────────────────────────────────────────────────────

    def _papers_for_download(self, selected_only: bool) -> List[Dict[str, Any]]:
        if selected_only:
            return [e["paper"] for e in self._checked_entries()]
        return [
            self.papers_table.item(r, 0).data(Qt.ItemDataRole.UserRole)["paper"]
            for r in range(self.papers_table.rowCount())
            if self.papers_table.item(r, 0)
        ]

    def _start_download(self, selected_only: bool):
        papers = self._papers_for_download(selected_only)
        if not papers:
            QMessageBox.information(self, "Nothing to download",
                                    "Select papers or load a list first.")
            return
        self._dl_progress.setMaximum(len(papers))
        self._dl_progress.setValue(0)
        self._dl_progress.setVisible(True)
        self._stop_dl_btn.setVisible(True)
        self._dl_status.setText(f"Starting download of {len(papers)} PDF(s)…")

        self._download_worker = BatchPdfDownloadWorker(papers)
        self._download_thread = QThread()
        self._download_worker.moveToThread(self._download_thread)
        self._download_thread.started.connect(self._download_worker.run)
        self._download_worker.progress.connect(self._on_dl_progress)
        self._download_worker.finished.connect(self._on_dl_finished)
        self._download_worker.stopped.connect(self._on_dl_stopped)
        self._download_worker.finished.connect(self._download_thread.quit)
        self._download_worker.stopped.connect(self._download_thread.quit)
        self._download_thread.start()

    def _stop_download(self):
        if self._download_worker:
            self._download_worker.stop()
        self._stop_dl_btn.setEnabled(False)

    def _on_dl_progress(self, done: int, total: int, title: str):
        self._dl_progress.setMaximum(total)
        self._dl_progress.setValue(done)
        self._dl_status.setText(f"Downloading {done + 1}/{total}: {title}…")

    def _on_dl_finished(self, succeeded: int, failed: int):
        self._dl_progress.setVisible(False)
        self._stop_dl_btn.setVisible(True)
        self._stop_dl_btn.setEnabled(True)
        self._stop_dl_btn.setVisible(False)
        msg = f"Done — {succeeded} downloaded"
        if failed:
            msg += f", {failed} failed (no direct PDF URL)"
        self._dl_status.setText(msg)

    def _on_dl_stopped(self):
        self._dl_progress.setVisible(False)
        self._stop_dl_btn.setVisible(False)
        self._stop_dl_btn.setEnabled(True)
        self._dl_status.setText("Download stopped.")


# ---------------------------------------------------------------------------
# Settings tab
# ---------------------------------------------------------------------------

SETTINGS_FILES = [
    ("sources_config.yaml", "Sources — which APIs are enabled and your contact email"),
    ("llm_config.yaml",     "LLM — provider, model, and token budget"),
]


class SettingsTab(QWidget):
    """Edit YAML configuration files. Changes take effect on next restart."""

    def __init__(self):
        super().__init__()
        self._paths = {label: Path(fname) for fname, label in SETTINGS_FILES}
        self._current_path: Optional[Path] = None
        self._init_ui()

    def _init_ui(self):
        layout = QVBoxLayout()

        top = QHBoxLayout()
        top.addWidget(QLabel("Edit:"))
        self.file_combo = QComboBox()
        for _, label in SETTINGS_FILES:
            self.file_combo.addItem(label)
        self.file_combo.currentIndexChanged.connect(self._load_file)
        top.addWidget(self.file_combo, 1)
        layout.addLayout(top)

        self.editor = QPlainTextEdit()
        self.editor.setFont(QFont("Menlo, Monaco, Courier New", 12))
        layout.addWidget(self.editor, 1)

        bottom = QHBoxLayout()
        self.status_label = QLabel("")
        bottom.addWidget(self.status_label, 1)
        reload_btn = QPushButton("Reload from disk")
        reload_btn.clicked.connect(self._load_file)
        save_btn = QPushButton("Save")
        save_btn.clicked.connect(self._save_file)
        bottom.addWidget(reload_btn)
        bottom.addWidget(save_btn)
        layout.addLayout(bottom)

        layout.addWidget(QLabel(
            "Changes take effect on next restart.",
            styleSheet="color: gray; font-size: 11px;"
        ))

        self.setLayout(layout)
        self._load_file()

    def _load_file(self):
        _, label = SETTINGS_FILES[self.file_combo.currentIndex()]
        path = self._paths[label]
        self._current_path = path
        try:
            self.editor.setPlainText(path.read_text())
            self.status_label.setText(f"Loaded {path.name}")
        except FileNotFoundError:
            self.editor.setPlainText("")
            self.status_label.setText(f"Not found: {path}")

    def _save_file(self):
        if self._current_path is None:
            return
        import yaml
        text = self.editor.toPlainText()
        try:
            yaml.safe_load(text)
        except yaml.YAMLError as e:
            QMessageBox.critical(self, "Invalid YAML", str(e))
            return
        try:
            self._current_path.write_text(text)
            self.status_label.setText(f"Saved {self._current_path.name} — restart to apply.")
        except OSError as e:
            QMessageBox.critical(self, "Save failed", str(e))


# ---------------------------------------------------------------------------

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Research Tool")
        self.setGeometry(100, 100, 1400, 850)

        self.db = Database()

        # Build orchestrator from config
        sources_config = load_sources_config()
        self.orchestrator = SourceOrchestrator(sources_config)

        self._show_startup_warnings(self.orchestrator.warnings)

        tabs = QTabWidget()

        self.search_tab  = SearchBrowseTab(self.db, self.orchestrator)
        self.filters_tab = FiltersTab(self.orchestrator)

        self.saved_refs_tab = SavedReferencesTab(self.db)

        # When filters are saved, refresh the search tab's list
        self.filters_tab.filters_changed.connect(self.search_tab.refresh_filters)
        # When a reference list is saved, refresh the Saved References tab
        self.search_tab.reference_list_saved.connect(self.saved_refs_tab.refresh_lists)

        self.settings_tab = SettingsTab()

        tabs.addTab(self.search_tab,     "Search & Browse")
        tabs.addTab(self.filters_tab,    "Filters")
        tabs.addTab(self.saved_refs_tab, "Saved References")
        tabs.addTab(self.settings_tab,   "Settings")

        self.setCentralWidget(tabs)
        logger.info("Application started with sources: %s",
                    self.orchestrator.get_enabled_sources())

    def _show_startup_warnings(self, warnings: list) -> None:
        for msg in warnings:
            logger.warning("Startup: %s", msg)
        if warnings:
            # Join all warnings so each is visible; individual lines go to the log.
            self.statusBar().showMessage("⚠ " + " | ".join(warnings), 0)

    def closeEvent(self, event):
        self.db.close()
        event.accept()


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
